# Перенос данных из старого Excel-отчёта ПТО в базу новой системы.
# python3 -I tools/import_xlsx.py <старый.xlsx> <папка_выхода>
#   → База_ПТО.xlsx (загрузить в Google Диск и открыть как Google Таблицу)
#   → data.json     (те же данные для проверки расчётов и тестовой страницы)
import sys, json, datetime, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

src, out = sys.argv[1], sys.argv[2]
os.makedirs(out, exist_ok=True)
wb = openpyxl.load_workbook(src)

def d(v):
    if isinstance(v, datetime.datetime): return v.strftime('%Y-%m-%d')
    if isinstance(v, datetime.date): return v.isoformat()
    return ''
def m(v):
    return d(v)[:7] if v else ''
def n(v):
    if v is None or v == '': return None
    try: return float(v)
    except Exception: return None
def s(v):
    return '' if v is None else str(v).strip()

objs = []
ws = wb['Объекты']
for r in range(5, 85):
    name = s(ws.cell(r, 2).value)
    if not name: continue
    c = lambda k: ws.cell(r, k).value
    objs.append({'id': 'o%02d' % (r - 4), 'name': name, 'resp': s(c(3)), 'pto': s(c(19)), 'status': s(c(4)) or 'В работе',
                 'sum': n(c(5)), 'done0': n(c(6)) or 0, 'ret': n(c(7)) if c(7) is not None else 0.05,
                 'payDays': int(n(c(8)) or 30), 'due': d(c(9)), 'sched': s(c(20)), 'note': s(c(10))})

ks = []
ws = wb['Журнал КС']
for r in range(6, 1006):
    c = lambda k: ws.cell(r, k).value
    if not (c(4) or c(7)): continue
    ks.append({'id': 'k%04d' % (r - 5), 'date': d(c(2)), 'month': m(c(3)) or m(c(2)), 'obj': s(c(4)), 'num': s(c(5)),
               'period': s(c(6)), 'ks2': n(c(7)), 'ks3': n(c(8)), 'agr': s(c(9)), 'sign': s(c(10)), 'pay': s(c(11)),
               'payDate': d(c(12)), 'signDate': d(c(20)), 'note': s(c(13))})

pays = []
ws = wb['Поступления']
for r in range(6, 506):
    c = lambda k: ws.cell(r, k).value
    if not (c(2) or c(3) or c(4)): continue
    pays.append({'id': 'p%04d' % (r - 5), 'date': d(c(2)), 'obj': s(c(3)), 'sum': n(c(4)), 'kind': s(c(5)), 'basis': s(c(6)), 'note': s(c(7))})

plan = []
ws = wb['План']
months = {col: m(ws.cell(4, col).value) for col in range(4, 28)}
for r in range(5, 85):
    on = objs and next((o['name'] for o in objs if o['id'] == 'o%02d' % (r - 4)), None)
    if not on: continue
    for col, mo in months.items():
        v = n(ws.cell(r, col).value)
        if v: plan.append({'id': on + '|' + mo, 'obj': on, 'month': mo, 'sum': v})

gar = []
ws = wb['Гарантии']
for r in range(4, 44):
    c = lambda k: ws.cell(r, k).value
    if not (c(2) or c(4)): continue
    gar.append({'id': 'g%03d' % (r - 3), 'obj': s(c(2)), 'due': d(c(3)), 'sum': n(c(4)), 'status': s(c(5)) or 'Не оплачена', 'party': s(c(6)), 'note': s(c(7))})

notes = []
ws = wb['По месяцам']
for r in range(42, 66):
    mo, txt = m(ws.cell(r, 1).value), s(ws.cell(r, 2).value)
    if mo and txt: notes.append({'id': mo, 'month': mo, 'text': txt})

ws = wb['Списки']
col = lambda k: [s(ws.cell(r, k).value) for r in range(4, 24) if s(ws.cell(r, k).value) and not s(ws.cell(r, k).value).startswith('Пустые')]
lists = {'resp': col(2), 'pto': col(14)}

data = {'objects': objs, 'ks': ks, 'pays': pays, 'plan': plan, 'gar': gar, 'notes': notes, 'lists': lists}
json.dump(data, open(os.path.join(out, 'data.json'), 'w'), ensure_ascii=False, indent=1)

# ---------- База для Google Таблицы: те же листы и столбцы, что в apps-script/Code.gs ----------
SCHEMA = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'src', 'schema.json')))
bw = openpyxl.Workbook(); bw.remove(bw.active)
HDR = PatternFill('solid', fgColor='2E671F')
for key, (title, cols) in SCHEMA.items():
    sh = bw.create_sheet(title)
    for j, (field, label, typ) in enumerate(cols, 1):
        cell = sh.cell(1, j, label); cell.font = Font(bold=True, color='FFFFFF'); cell.fill = HDR
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        sh.column_dimensions[openpyxl.utils.get_column_letter(j)].width = {'s': 24, 'd': 12, 'm': 12, 'n': 16, 'p': 9}[typ] + (16 if field in ('name', 'num', 'note', 'obj', 'text') else 0)
    sh.freeze_panes = 'A2'
    rows = {'O': objs, 'K': ks, 'P': pays, 'L': plan, 'G': gar, 'N': notes}.get(key)
    if key == 'S':
        mx = max(len(lists['resp']), len(lists['pto']))
        rows = [{'resp': (lists['resp'] + [''] * mx)[i], 'pto': (lists['pto'] + [''] * mx)[i]} for i in range(mx)]
    for i, row in enumerate(rows or [], 2):
        for j, (field, label, typ) in enumerate(cols, 1):
            v = row.get(field)
            if v in (None, ''): continue
            if typ == 'd': v = datetime.datetime.strptime(v, '%Y-%m-%d')
            if typ == 'm': v = datetime.datetime.strptime(v + '-01', '%Y-%m-%d')
            cell = sh.cell(i, j, v)
            cell.number_format = {'d': 'DD.MM.YYYY', 'm': 'MM.YYYY', 'n': '#,##0.00', 'p': '0.0%'}.get(typ, 'General')
bw.save(os.path.join(out, 'База_ПТО.xlsx'))
print('объектов', len(objs), 'КС', len(ks), 'поступлений', len(pays), 'план', len(plan), 'гарантий', len(gar), 'комментариев', len(notes))
